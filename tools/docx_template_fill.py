#!/usr/bin/env python3
"""Fill an existing DOCX/DOTX template with an already-rendered DOCX body.

The template is the base OPC package.  Only ``word/document.xml`` and the
parts that are required by the inserted body are changed, so cover pages,
content controls, headers/footers, form fields, tables, styles, and custom XML
from the template remain intact.

Typical usage::

    python tools/docx_template_fill.py \
        --template school-template.docx \
        --content-docx rendered-body.docx \
        --output final.docx \
        --map _template_map.json

``body_anchor_para_idx`` in the optional map follows python-docx's
``Document.paragraphs`` convention (zero-based, direct body paragraphs).  A
map can explicitly set ``paragraph_index_base`` to ``1`` when its analyzer
uses one-based indices.  Without a map, common BODY/正文 placeholders are
detected conservatively.
"""

from __future__ import annotations

import argparse
import copy
import json
import os
import posixpath
import re
import sys
import tempfile
import unicodedata
import zipfile
from collections import OrderedDict
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Iterable, Iterator, Mapping, MutableMapping, Sequence
from urllib.parse import quote, unquote, urlsplit

from lxml import etree


W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
CT_NS = "http://schemas.openxmlformats.org/package/2006/content-types"
WP_NS = "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
PIC_NS = "http://schemas.openxmlformats.org/drawingml/2006/picture"

DOCUMENT_PART = "word/document.xml"
DOCUMENT_RELS_PART = "word/_rels/document.xml.rels"
CONTENT_TYPES_PART = "[Content_Types].xml"

DOCX_MAIN_CONTENT_TYPE = (
    "application/vnd.openxmlformats-officedocument."
    "wordprocessingml.document.main+xml"
)
STYLES_CONTENT_TYPE = (
    "application/vnd.openxmlformats-officedocument."
    "wordprocessingml.styles+xml"
)
NUMBERING_CONTENT_TYPE = (
    "application/vnd.openxmlformats-officedocument."
    "wordprocessingml.numbering+xml"
)
FOOTNOTES_CONTENT_TYPE = (
    "application/vnd.openxmlformats-officedocument."
    "wordprocessingml.footnotes+xml"
)
ENDNOTES_CONTENT_TYPE = (
    "application/vnd.openxmlformats-officedocument."
    "wordprocessingml.endnotes+xml"
)
COMMENTS_CONTENT_TYPE = (
    "application/vnd.openxmlformats-officedocument."
    "wordprocessingml.comments+xml"
)

RELATIONSHIP_TYPE_BASE = (
    "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
)

_XML_PARSER = etree.XMLParser(
    resolve_entities=False,
    no_network=True,
    recover=False,
    remove_blank_text=False,
    huge_tree=False,
)


class TemplateFillError(RuntimeError):
    """Raised for a malformed input package or an unusable template map."""


@dataclass(frozen=True)
class FillReport:
    template: str
    content_docx: str
    output: str
    anchor_source: str
    anchor_text: str
    inserted_elements: int
    imported_relationships: int
    imported_parts: int
    imported_styles: int
    imported_numbering_instances: int


def _qn(namespace: str, local_name: str) -> str:
    return f"{{{namespace}}}{local_name}"


def _namespace(tag: str) -> str:
    qname = etree.QName(tag)
    return qname.namespace or ""


def _local_name(tag: str) -> str:
    return etree.QName(tag).localname


def _parse_xml(data: bytes, part_name: str) -> etree._Element:
    try:
        return etree.fromstring(data, parser=_XML_PARSER)
    except (etree.XMLSyntaxError, ValueError) as exc:
        raise TemplateFillError(f"Invalid XML part {part_name!r}: {exc}") from exc


def _serialize_xml(root: etree._Element) -> bytes:
    return etree.tostring(
        root,
        encoding="UTF-8",
        xml_declaration=True,
        standalone=True,
    )


def _rels_part_name(part_name: str) -> str:
    directory, basename = posixpath.split(part_name)
    return posixpath.join(directory, "_rels", f"{basename}.rels")


def _safe_part_name(value: str) -> str:
    value = value.replace("\\", "/").lstrip("/")
    value = posixpath.normpath(value)
    if value in {"", ".", ".."} or value.startswith("../"):
        raise TemplateFillError(f"Unsafe OPC part path: {value!r}")
    return value


def _resolve_relationship_target(source_part: str, target: str) -> str:
    parsed = urlsplit(target)
    raw_path = unquote(parsed.path)
    if not raw_path:
        raise TemplateFillError(
            f"Empty internal relationship target in {source_part!r}: {target!r}"
        )
    if raw_path.startswith("/"):
        return _safe_part_name(raw_path)
    return _safe_part_name(posixpath.join(posixpath.dirname(source_part), raw_path))


def _relative_relationship_target(source_part: str, target_part: str) -> str:
    relative = posixpath.relpath(target_part, posixpath.dirname(source_part) or ".")
    # OPC relationship targets are URI references.  Keep path separators and
    # the common unreserved filename characters readable.
    return quote(relative, safe="/._-~")


class _Package:
    """An in-memory OPC ZIP package that preserves untouched ZIP metadata."""

    def __init__(self, path: Path) -> None:
        self.path = path
        self.parts: MutableMapping[str, bytes] = OrderedDict()
        self.infos: dict[str, zipfile.ZipInfo] = {}
        try:
            with zipfile.ZipFile(path, "r") as archive:
                bad_member = archive.testzip()
                if bad_member:
                    raise TemplateFillError(
                        f"Corrupt ZIP member {bad_member!r} in {str(path)!r}"
                    )
                for info in archive.infolist():
                    if info.is_dir():
                        continue
                    name = _safe_part_name(info.filename)
                    if name in self.parts:
                        raise TemplateFillError(
                            f"Duplicate OPC part {name!r} in {str(path)!r}"
                        )
                    self.parts[name] = archive.read(info)
                    self.infos[name] = info
        except (OSError, zipfile.BadZipFile, RuntimeError) as exc:
            if isinstance(exc, TemplateFillError):
                raise
            raise TemplateFillError(f"Cannot read DOCX package {str(path)!r}: {exc}") from exc

        for required in (CONTENT_TYPES_PART, DOCUMENT_PART):
            if required not in self.parts:
                raise TemplateFillError(
                    f"{str(path)!r} is not a Word package: missing {required!r}"
                )

    def write(self, output: Path) -> None:
        output.parent.mkdir(parents=True, exist_ok=True)
        fd, temp_name = tempfile.mkstemp(
            prefix=f".{output.name}.", suffix=".tmp", dir=str(output.parent)
        )
        os.close(fd)
        temp_path = Path(temp_name)
        try:
            with zipfile.ZipFile(
                temp_path,
                "w",
                compression=zipfile.ZIP_DEFLATED,
                compresslevel=6,
                allowZip64=True,
            ) as archive:
                for name, data in self.parts.items():
                    original = self.infos.get(name)
                    if original is not None:
                        info = copy.copy(original)
                        info.filename = name
                        archive.writestr(info, data)
                    else:
                        info = zipfile.ZipInfo(name, date_time=(1980, 1, 1, 0, 0, 0))
                        info.compress_type = zipfile.ZIP_DEFLATED
                        info.external_attr = 0o600 << 16
                        archive.writestr(info, data)
            with zipfile.ZipFile(temp_path, "r") as archive:
                bad_member = archive.testzip()
                if bad_member:
                    raise TemplateFillError(
                        f"Generated DOCX contains a corrupt member: {bad_member!r}"
                    )
            os.replace(temp_path, output)
        finally:
            try:
                temp_path.unlink(missing_ok=True)
            except OSError:
                pass


class _ContentTypes:
    def __init__(self, package: _Package) -> None:
        self.package = package
        self.root = _parse_xml(package.parts[CONTENT_TYPES_PART], CONTENT_TYPES_PART)
        self.ns = _namespace(self.root.tag) or CT_NS

    def _defaults(self) -> dict[str, etree._Element]:
        result: dict[str, etree._Element] = {}
        for child in self.root:
            if _local_name(child.tag) == "Default":
                result[child.get("Extension", "").lower()] = child
        return result

    def _overrides(self) -> dict[str, etree._Element]:
        result: dict[str, etree._Element] = {}
        for child in self.root:
            if _local_name(child.tag) == "Override":
                result[child.get("PartName", "").lstrip("/")] = child
        return result

    def get(self, part_name: str) -> str | None:
        part_name = part_name.lstrip("/")
        override = self._overrides().get(part_name)
        if override is not None:
            return override.get("ContentType")
        extension = posixpath.splitext(part_name)[1].lstrip(".").lower()
        default = self._defaults().get(extension)
        return default.get("ContentType") if default is not None else None

    def ensure(self, part_name: str, content_type: str | None) -> None:
        if not content_type:
            return
        part_name = part_name.lstrip("/")
        overrides = self._overrides()
        existing_override = overrides.get(part_name)
        if existing_override is not None:
            # A template-owned part wins.  For newly copied parts, the caller
            # allocated a fresh name, so this branch normally means identical
            # content types.
            return

        extension = posixpath.splitext(part_name)[1].lstrip(".").lower()
        existing_default = self._defaults().get(extension)
        if existing_default is None and extension:
            default = etree.Element(_qn(self.ns, "Default"))
            default.set("Extension", extension)
            default.set("ContentType", content_type)
            first_override = next(
                (
                    index
                    for index, child in enumerate(self.root)
                    if _local_name(child.tag) == "Override"
                ),
                len(self.root),
            )
            self.root.insert(first_override, default)
            return
        if existing_default is not None and existing_default.get("ContentType") == content_type:
            return

        override = etree.Element(_qn(self.ns, "Override"))
        override.set("PartName", f"/{part_name}")
        override.set("ContentType", content_type)
        self.root.append(override)

    def force_docx_main_document(self) -> None:
        override = self._overrides().get(DOCUMENT_PART)
        if override is None:
            override = etree.Element(_qn(self.ns, "Override"))
            override.set("PartName", f"/{DOCUMENT_PART}")
            self.root.append(override)
        override.set("ContentType", DOCX_MAIN_CONTENT_TYPE)

    def flush(self) -> None:
        self.package.parts[CONTENT_TYPES_PART] = _serialize_xml(self.root)


class _RelationshipRoot:
    def __init__(self, root: etree._Element | None = None) -> None:
        if root is None:
            root = etree.Element(_qn(REL_NS, "Relationships"), nsmap={None: REL_NS})
        self.root = root
        self.ns = _namespace(root.tag) or REL_NS

    def relationships(self) -> list[etree._Element]:
        return [child for child in self.root if _local_name(child.tag) == "Relationship"]

    def by_id(self) -> dict[str, etree._Element]:
        return {rel.get("Id", ""): rel for rel in self.relationships()}

    def allocate_id(self) -> str:
        used = set(self.by_id())
        highest = 0
        for value in used:
            match = re.fullmatch(r"rId(\d+)", value)
            if match:
                highest = max(highest, int(match.group(1)))
        candidate = highest + 1
        while f"rId{candidate}" in used:
            candidate += 1
        return f"rId{candidate}"

    def add_from(
        self,
        source: etree._Element,
        relationship_id: str,
        target: str,
    ) -> etree._Element:
        created = etree.Element(_qn(self.ns, "Relationship"))
        for key, value in source.attrib.items():
            if key not in {"Id", "Target"}:
                created.set(key, value)
        created.set("Id", relationship_id)
        created.set("Target", target)
        self.root.append(created)
        return created

    def add(self, relationship_id: str, rel_type: str, target: str) -> etree._Element:
        created = etree.Element(_qn(self.ns, "Relationship"))
        created.set("Id", relationship_id)
        created.set("Type", rel_type)
        created.set("Target", target)
        self.root.append(created)
        return created


class _OpcMerger:
    """Copy relationship-reachable parts from content into the template."""

    def __init__(self, target: _Package, source: _Package) -> None:
        self.target = target
        self.source = source
        self.target_content_types = _ContentTypes(target)
        self.source_content_types = _ContentTypes(source)
        self.part_map: dict[str, str] = {}
        self._relationship_roots: dict[str, _RelationshipRoot] = {}
        self._source_relationship_roots: dict[str, _RelationshipRoot] = {}
        self.imported_relationships = 0
        self.imported_parts = 0

    def _load_target_relationships(self, source_part: str) -> _RelationshipRoot:
        rels_name = _rels_part_name(source_part)
        cached = self._relationship_roots.get(rels_name)
        if cached is not None:
            return cached
        if rels_name in self.target.parts:
            root = _RelationshipRoot(
                _parse_xml(self.target.parts[rels_name], rels_name)
            )
        else:
            root = _RelationshipRoot()
        self._relationship_roots[rels_name] = root
        return root

    def _load_source_relationships(self, source_part: str) -> _RelationshipRoot | None:
        rels_name = _rels_part_name(source_part)
        cached = self._source_relationship_roots.get(rels_name)
        if cached is not None:
            return cached
        data = self.source.parts.get(rels_name)
        if data is None:
            return None
        root = _RelationshipRoot(_parse_xml(data, rels_name))
        self._source_relationship_roots[rels_name] = root
        return root

    def source_related_part(self, rel_type_suffix: str, default: str) -> str:
        rels = self._load_source_relationships(DOCUMENT_PART)
        if rels is not None:
            for rel in rels.relationships():
                if (rel.get("Type") or "").endswith(rel_type_suffix):
                    if (rel.get("TargetMode") or "").lower() != "external":
                        return _resolve_relationship_target(
                            DOCUMENT_PART, rel.get("Target", "")
                        )
        return default

    def target_related_part(self, rel_type_suffix: str, default: str) -> str:
        rels = self._load_target_relationships(DOCUMENT_PART)
        for rel in rels.relationships():
            if (rel.get("Type") or "").endswith(rel_type_suffix):
                if (rel.get("TargetMode") or "").lower() != "external":
                    return _resolve_relationship_target(
                        DOCUMENT_PART, rel.get("Target", "")
                    )
        return default

    def ensure_document_relationship(
        self,
        target_part: str,
        rel_type_suffix: str,
    ) -> None:
        relationships = self._load_target_relationships(DOCUMENT_PART)
        for rel in relationships.relationships():
            if (rel.get("Type") or "").endswith(rel_type_suffix):
                resolved = _resolve_relationship_target(
                    DOCUMENT_PART, rel.get("Target", "")
                )
                if resolved == target_part:
                    return

        source_relationships = self._load_source_relationships(DOCUMENT_PART)
        relationship_type = f"{RELATIONSHIP_TYPE_BASE}{rel_type_suffix}"
        if source_relationships is not None:
            for rel in source_relationships.relationships():
                if (rel.get("Type") or "").endswith(rel_type_suffix):
                    relationship_type = rel.get("Type") or relationship_type
                    break
        relationships.add(
            relationships.allocate_id(),
            relationship_type,
            _relative_relationship_target(DOCUMENT_PART, target_part),
        )

    def _allocate_target_part(self, source_part: str) -> str:
        if source_part not in self.target.parts and source_part not in self.part_map.values():
            return source_part
        directory, basename = posixpath.split(source_part)
        stem, extension = posixpath.splitext(basename)
        index = 1
        while True:
            candidate = posixpath.join(
                directory, f"content_{index}_{stem}{extension}"
            )
            if candidate not in self.target.parts and candidate not in self.part_map.values():
                return candidate
            index += 1

    def copy_part_recursive(self, source_part: str) -> str:
        source_part = _safe_part_name(source_part)
        mapped = self.part_map.get(source_part)
        if mapped is not None:
            return mapped
        if source_part not in self.source.parts:
            # A reference-doc render can intentionally retain a relationship to
            # a template-owned part.  Reuse it only when it really exists.
            if source_part in self.target.parts:
                self.part_map[source_part] = source_part
                return source_part
            raise TemplateFillError(
                f"Content relationship points to missing OPC part {source_part!r}"
            )

        target_part = self._allocate_target_part(source_part)
        self.part_map[source_part] = target_part  # break recursive cycles
        self.target.parts[target_part] = self.source.parts[source_part]
        self.target_content_types.ensure(
            target_part, self.source_content_types.get(source_part)
        )
        self.imported_parts += 1

        source_rels_name = _rels_part_name(source_part)
        if source_rels_name in self.source.parts:
            source_rels = _RelationshipRoot(
                _parse_xml(self.source.parts[source_rels_name], source_rels_name)
            )
            copied_root = copy.deepcopy(source_rels.root)
            copied_relationships = _RelationshipRoot(copied_root)
            for relationship in copied_relationships.relationships():
                if (relationship.get("TargetMode") or "").lower() == "external":
                    continue
                dependency = _resolve_relationship_target(
                    source_part, relationship.get("Target", "")
                )
                copied_dependency = self.copy_part_recursive(dependency)
                relationship.set(
                    "Target",
                    _relative_relationship_target(target_part, copied_dependency),
                )
            target_rels_name = _rels_part_name(target_part)
            self.target.parts[target_rels_name] = _serialize_xml(copied_root)
            self.target_content_types.ensure(
                target_rels_name, self.source_content_types.get(source_rels_name)
            )

        return target_part

    @staticmethod
    def _relationship_attributes(
        nodes: Iterable[etree._Element],
    ) -> Iterator[tuple[etree._Element, str, str]]:
        for node in nodes:
            for element in node.iter():
                for attribute_name, value in list(element.attrib.items()):
                    qname = etree.QName(attribute_name)
                    namespace = qname.namespace or ""
                    if (
                        qname.localname in {"id", "embed", "link"}
                        and namespace.endswith("/relationships")
                        and value
                    ):
                        yield element, attribute_name, value

    def merge_relationships_for_nodes(
        self,
        source_part: str,
        target_part: str,
        nodes: Sequence[etree._Element],
    ) -> int:
        attributes = list(self._relationship_attributes(nodes))
        if not attributes:
            return 0
        source_relationships = self._load_source_relationships(source_part)
        target_relationships = self._load_target_relationships(target_part)
        source_by_id = source_relationships.by_id() if source_relationships else {}
        target_by_id = target_relationships.by_id()
        id_map: dict[str, str] = {}

        for _, _, old_id in attributes:
            if old_id in id_map:
                continue
            source_relationship = source_by_id.get(old_id)
            if source_relationship is None:
                # If the content was rendered against this exact reference
                # document, a template-owned relationship with the same rId is
                # a safe and useful fallback.
                if old_id in target_by_id:
                    id_map[old_id] = old_id
                    continue
                raise TemplateFillError(
                    f"Relationship {old_id!r} used in {source_part!r} is missing"
                )

            new_id = target_relationships.allocate_id()
            if (source_relationship.get("TargetMode") or "").lower() == "external":
                new_target = source_relationship.get("Target", "")
            else:
                source_dependency = _resolve_relationship_target(
                    source_part, source_relationship.get("Target", "")
                )
                target_dependency = self.copy_part_recursive(source_dependency)
                new_target = _relative_relationship_target(
                    target_part, target_dependency
                )
            target_relationships.add_from(source_relationship, new_id, new_target)
            id_map[old_id] = new_id
            self.imported_relationships += 1

        for element, attribute_name, old_id in attributes:
            element.set(attribute_name, id_map[old_id])
        return len(id_map)

    def flush_relationships(self) -> None:
        for part_name, relationships in self._relationship_roots.items():
            self.target.parts[part_name] = _serialize_xml(relationships.root)
            self.target_content_types.ensure(
                part_name,
                self.target_content_types.get(part_name)
                or "application/vnd.openxmlformats-package.relationships+xml",
            )


def _body_element(document_root: etree._Element, word_ns: str) -> etree._Element:
    body = document_root.find(_qn(word_ns, "body"))
    if body is None:
        raise TemplateFillError("word/document.xml has no w:body element")
    return body


def _paragraph_text(paragraph: etree._Element, word_ns: str) -> str:
    text_tags = {
        _qn(word_ns, "t"),
        _qn(word_ns, "delText"),
        _qn(word_ns, "instrText"),
    }
    return "".join(
        element.text or "" for element in paragraph.iter() if element.tag in text_tags
    ).strip()


def _normalized_placeholder_text(value: str) -> str:
    value = unicodedata.normalize("NFKC", value)
    value = value.replace("\u200b", "").replace("\ufeff", "")
    return re.sub(r"\s+", "", value).strip().upper()


def _placeholder_score(value: str) -> int:
    text = _normalized_placeholder_text(value)
    if not text:
        return 0
    exact = {
        "{{BODY}}",
        "{{CONTENT}}",
        "{{MAIN_CONTENT}}",
        "{{ARTICLE_BODY}}",
        "{{正文}}",
        "[[BODY]]",
        "[[CONTENT]]",
        "[[正文]]",
        "[BODY]",
        "[CONTENT]",
        "[正文]",
        "<BODY>",
        "<CONTENT>",
        "<正文>",
        "${BODY}",
        "${CONTENT}",
        "__BODY__",
        "__CONTENT__",
        "BODY_PLACEHOLDER",
        "CONTENT_PLACEHOLDER",
        "正文占位符",
        "正文内容占位符",
    }
    if text in exact:
        return 120
    if re.fullmatch(r"[<{\[(_$%-]*BODY[>}\])_%.-]*", text):
        return 115
    if re.fullmatch(r"[<{\[(_$%-]*(MAIN)?CONTENT[>}\])_%.-]*", text):
        return 112
    if re.search(r"(在?此处|请)(插入|填写|填入|替换|粘贴).{0,4}(论文)?正文", text):
        return 108
    if re.search(r"(INSERT|PLACE|PUT|WRITE).{0,12}(BODY|MAINCONTENT).*HERE", text):
        return 105
    if text.startswith("正文内容(") or text.startswith("正文内容（"):
        return 100
    if text in {"正文内容", "论文正文", "MAINCONTENT", "ARTICLEBODY"}:
        return 92
    if text in {"正文", "BODY"}:
        return 80
    return 0


def _load_template_map(map_path: Path | None) -> Mapping[str, object]:
    if map_path is None:
        return {}
    try:
        data = json.loads(map_path.read_text(encoding="utf-8-sig"))
    except (OSError, UnicodeError, json.JSONDecodeError) as exc:
        raise TemplateFillError(f"Cannot read template map {str(map_path)!r}: {exc}") from exc
    if not isinstance(data, dict):
        raise TemplateFillError("Template map must be a JSON object")
    return data


def _map_anchor_index(template_map: Mapping[str, object]) -> int | None:
    value: object | None = None
    for key in ("body_anchor_para_idx", "body_anchor_index", "anchor_para_idx"):
        if key in template_map:
            value = template_map[key]
            break
    anchors = template_map.get("anchors")
    if value is None and isinstance(anchors, Mapping):
        value = anchors.get("body")
        if isinstance(value, Mapping):
            value = value.get("paragraph_index", value.get("para_idx"))
    if value is None:
        return None
    if isinstance(value, bool):
        raise TemplateFillError("body_anchor_para_idx must be an integer, not boolean")
    try:
        parsed = int(value)
    except (TypeError, ValueError) as exc:
        raise TemplateFillError("body_anchor_para_idx must be an integer") from exc
    if parsed < 0:
        raise TemplateFillError("body_anchor_para_idx cannot be negative")
    return parsed


def _anchor_from_map(
    body: etree._Element,
    word_ns: str,
    template_map: Mapping[str, object],
) -> tuple[etree._Element | None, str | None]:
    raw_index = _map_anchor_index(template_map)
    if raw_index is None:
        return None, None

    scope = str(template_map.get("paragraph_scope", "direct")).strip().lower()
    direct = [child for child in body if child.tag == _qn(word_ns, "p")]
    all_paragraphs = list(body.iter(_qn(word_ns, "p")))
    paragraphs = all_paragraphs if scope in {"all", "descendants"} else direct

    index_base = template_map.get(
        "paragraph_index_base", template_map.get("index_base")
    )
    if index_base is not None:
        try:
            base = int(index_base)
        except (TypeError, ValueError) as exc:
            raise TemplateFillError("paragraph_index_base must be 0 or 1") from exc
        if base not in {0, 1}:
            raise TemplateFillError("paragraph_index_base must be 0 or 1")
        index = raw_index - base
        if not 0 <= index < len(paragraphs):
            raise TemplateFillError(
                f"body_anchor_para_idx={raw_index} is outside the {len(paragraphs)} "
                f"template paragraphs (index base {base})"
            )
        return paragraphs[index], f"map:index={raw_index},base={base},scope={scope}"

    # Default is python-docx's zero-based convention.  For compatibility with
    # analyzers that print human-facing one-based paragraph numbers, prefer
    # index-1 only when it is recognizably a stronger BODY placeholder.
    zero_candidate = paragraphs[raw_index] if raw_index < len(paragraphs) else None
    one_candidate = (
        paragraphs[raw_index - 1]
        if raw_index > 0 and raw_index - 1 < len(paragraphs)
        else None
    )
    zero_score = (
        _placeholder_score(_paragraph_text(zero_candidate, word_ns))
        if zero_candidate is not None
        else 0
    )
    one_score = (
        _placeholder_score(_paragraph_text(one_candidate, word_ns))
        if one_candidate is not None
        else 0
    )
    if one_candidate is not None and one_score >= 80 and one_score > zero_score:
        return one_candidate, f"map:index={raw_index},auto-base=1,scope={scope}"
    if zero_candidate is not None:
        return zero_candidate, f"map:index={raw_index},base=0,scope={scope}"
    raise TemplateFillError(
        f"body_anchor_para_idx={raw_index} is outside the {len(paragraphs)} "
        "template paragraphs"
    )


def _find_placeholder(
    body: etree._Element, word_ns: str
) -> tuple[etree._Element | None, str | None]:
    best: tuple[int, int, etree._Element] | None = None
    for index, paragraph in enumerate(body.iter(_qn(word_ns, "p"))):
        score = _placeholder_score(_paragraph_text(paragraph, word_ns))
        candidate = (score, -index, paragraph)
        if score and (best is None or candidate[:2] > best[:2]):
            best = candidate
    if best is None:
        return None, None
    return best[2], f"placeholder:score={best[0]}"


def _content_body_nodes(
    content_body: etree._Element, word_ns: str
) -> list[etree._Element]:
    result: list[etree._Element] = []
    sect_pr_tag = _qn(word_ns, "sectPr")
    paragraph_tag = _qn(word_ns, "p")
    for source_child in content_body:
        if source_child.tag == sect_pr_tag:
            continue
        copied = copy.deepcopy(source_child)
        section_properties = list(copied.iter(sect_pr_tag))
        had_section_properties = bool(section_properties)
        for section_property in section_properties:
            parent = section_property.getparent()
            if parent is not None:
                parent.remove(section_property)
        # Pandoc/python-docx often leave a final paragraph that contains only a
        # section break.  Do not turn that into a spurious blank body paragraph.
        if (
            copied.tag == paragraph_tag
            and had_section_properties
            and not _paragraph_text(copied, word_ns)
            and not any(
                _local_name(element.tag) in {"drawing", "pict", "object", "altChunk"}
                for element in copied.iter()
            )
        ):
            continue
        result.append(copied)
    if not result:
        raise TemplateFillError("The content DOCX has no body elements to insert")
    return result


def _referenced_style_ids(
    nodes: Iterable[etree._Element], word_ns: str
) -> set[str]:
    style_tags = {
        _qn(word_ns, "pStyle"),
        _qn(word_ns, "rStyle"),
        _qn(word_ns, "tblStyle"),
    }
    value_attr = _qn(word_ns, "val")
    result: set[str] = set()
    for node in nodes:
        for element in node.iter():
            if element.tag in style_tags:
                value = element.get(value_attr)
                if value:
                    result.add(value)
    return result


def _insert_before_ext_list(root: etree._Element, element: etree._Element) -> None:
    for index, child in enumerate(root):
        if _local_name(child.tag) in {"extLst", "numIdMac"}:
            root.insert(index, element)
            return
    root.append(element)


def _merge_styles(
    merger: _OpcMerger,
    nodes: Sequence[etree._Element],
    word_ns: str,
) -> tuple[etree._Element | None, list[etree._Element]]:
    source_part = merger.source_related_part("/styles", "word/styles.xml")
    if source_part not in merger.source.parts:
        return None, []
    target_part = merger.target_related_part("/styles", "word/styles.xml")
    source_root = _parse_xml(merger.source.parts[source_part], source_part)

    if target_part in merger.target.parts:
        target_root = _parse_xml(merger.target.parts[target_part], target_part)
    else:
        target_root = etree.Element(
            _qn(word_ns, "styles"), nsmap={"w": word_ns, "r": R_NS}
        )
        # Copy defaults/latent-style metadata only when the template has no
        # styles part at all.  Existing template defaults must always win.
        for child in source_root:
            if _local_name(child.tag) in {"docDefaults", "latentStyles"}:
                target_root.append(copy.deepcopy(child))

    style_id_attr = _qn(word_ns, "styleId")
    value_attr = _qn(word_ns, "val")
    target_styles = {
        style.get(style_id_attr): style
        for style in target_root
        if _local_name(style.tag) == "style" and style.get(style_id_attr)
    }
    source_styles = {
        style.get(style_id_attr): style
        for style in source_root
        if _local_name(style.tag) == "style" and style.get(style_id_attr)
    }
    wanted = list(_referenced_style_ids(nodes, word_ns))
    copied_styles: list[etree._Element] = []
    visited: set[str] = set()

    def copy_style(style_id: str) -> None:
        if style_id in visited:
            return
        visited.add(style_id)
        if style_id in target_styles:
            return
        source_style = source_styles.get(style_id)
        if source_style is None:
            return
        for dependency_name in ("basedOn", "next", "link"):
            dependency = source_style.find(_qn(word_ns, dependency_name))
            if dependency is not None:
                dependency_id = dependency.get(value_attr)
                if dependency_id:
                    copy_style(dependency_id)
        copied = copy.deepcopy(source_style)
        _insert_before_ext_list(target_root, copied)
        target_styles[style_id] = copied
        copied_styles.append(copied)

    for wanted_style in wanted:
        copy_style(wanted_style)

    merger.target.parts[target_part] = _serialize_xml(target_root)
    merger.target_content_types.ensure(
        target_part,
        merger.source_content_types.get(source_part) or STYLES_CONTENT_TYPE,
    )
    merger.ensure_document_relationship(target_part, "/styles")
    return target_root, copied_styles


def _integer_attribute(element: etree._Element, attribute: str) -> int | None:
    value = element.get(attribute)
    if value is None:
        return None
    try:
        return int(value)
    except ValueError:
        return None


def _merge_numbering(
    merger: _OpcMerger,
    rewrite_nodes: Sequence[etree._Element],
    word_ns: str,
) -> tuple[etree._Element | None, int]:
    num_id_tag = _qn(word_ns, "numId")
    value_attr = _qn(word_ns, "val")
    used_ids: set[int] = set()
    for node in rewrite_nodes:
        for num_id in node.iter(num_id_tag):
            value = _integer_attribute(num_id, value_attr)
            if value is not None and value > 0:
                used_ids.add(value)
    if not used_ids:
        return None, 0

    source_part = merger.source_related_part("/numbering", "word/numbering.xml")
    target_part = merger.target_related_part("/numbering", "word/numbering.xml")
    source_data = merger.source.parts.get(source_part)
    if source_data is None:
        # Reference-document renders may use numbering already defined by the
        # template.  Verify every referenced instance exists before reusing it.
        target_data = merger.target.parts.get(target_part)
        if target_data is None:
            raise TemplateFillError("Content uses numbering but has no numbering.xml")
        target_root = _parse_xml(target_data, target_part)
        target_ids = {
            _integer_attribute(element, _qn(word_ns, "numId"))
            for element in target_root
            if _local_name(element.tag) == "num"
        }
        missing = used_ids - {value for value in target_ids if value is not None}
        if missing:
            raise TemplateFillError(
                f"Content references missing numbering instances: {sorted(missing)}"
            )
        return target_root, 0

    source_root = _parse_xml(source_data, source_part)
    if target_part in merger.target.parts:
        target_root = _parse_xml(merger.target.parts[target_part], target_part)
    else:
        target_root = etree.Element(
            _qn(word_ns, "numbering"), nsmap={"w": word_ns, "r": R_NS}
        )

    num_id_attr = _qn(word_ns, "numId")
    abstract_id_attr = _qn(word_ns, "abstractNumId")
    source_nums = {
        _integer_attribute(element, num_id_attr): element
        for element in source_root
        if _local_name(element.tag) == "num"
    }
    source_abstracts = {
        _integer_attribute(element, abstract_id_attr): element
        for element in source_root
        if _local_name(element.tag) == "abstractNum"
    }
    target_num_ids = {
        value
        for element in target_root
        if _local_name(element.tag) == "num"
        for value in [_integer_attribute(element, num_id_attr)]
        if value is not None
    }
    target_abstract_ids = {
        value
        for element in target_root
        if _local_name(element.tag) == "abstractNum"
        for value in [_integer_attribute(element, abstract_id_attr)]
        if value is not None
    }
    next_num_id = max(target_num_ids or {0}) + 1
    next_abstract_id = max(target_abstract_ids or {-1}) + 1
    num_map: dict[int, int] = {}
    abstract_map: dict[int, int] = {}
    copied_abstracts: list[etree._Element] = []
    copied_nums: list[etree._Element] = []

    for old_num_id in sorted(used_ids):
        source_num = source_nums.get(old_num_id)
        if source_num is None:
            if old_num_id in target_num_ids:
                num_map[old_num_id] = old_num_id
                continue
            raise TemplateFillError(
                f"Content numbering.xml has no w:num for numId={old_num_id}"
            )
        abstract_reference = source_num.find(_qn(word_ns, "abstractNumId"))
        old_abstract_id = (
            _integer_attribute(abstract_reference, value_attr)
            if abstract_reference is not None
            else None
        )
        if old_abstract_id is None or old_abstract_id not in source_abstracts:
            raise TemplateFillError(
                f"Content numbering numId={old_num_id} has no abstract definition"
            )
        if old_abstract_id not in abstract_map:
            while next_abstract_id in target_abstract_ids:
                next_abstract_id += 1
            new_abstract_id = next_abstract_id
            next_abstract_id += 1
            target_abstract_ids.add(new_abstract_id)
            copied_abstract = copy.deepcopy(source_abstracts[old_abstract_id])
            copied_abstract.set(abstract_id_attr, str(new_abstract_id))
            copied_abstracts.append(copied_abstract)
            abstract_map[old_abstract_id] = new_abstract_id

        while next_num_id in target_num_ids:
            next_num_id += 1
        new_num_id = next_num_id
        next_num_id += 1
        target_num_ids.add(new_num_id)
        copied_num = copy.deepcopy(source_num)
        copied_num.set(num_id_attr, str(new_num_id))
        copied_reference = copied_num.find(_qn(word_ns, "abstractNumId"))
        if copied_reference is not None:
            copied_reference.set(value_attr, str(abstract_map[old_abstract_id]))
        copied_nums.append(copied_num)
        num_map[old_num_id] = new_num_id

    # Keep numbering schema order: abstractNum entries precede num entries.
    first_num_index = next(
        (
            index
            for index, child in enumerate(target_root)
            if _local_name(child.tag) in {"num", "numIdMac"}
        ),
        len(target_root),
    )
    for copied_abstract in copied_abstracts:
        target_root.insert(first_num_index, copied_abstract)
        first_num_index += 1
    num_id_mac_index = next(
        (
            index
            for index, child in enumerate(target_root)
            if _local_name(child.tag) == "numIdMac"
        ),
        len(target_root),
    )
    for copied_num in copied_nums:
        target_root.insert(num_id_mac_index, copied_num)
        num_id_mac_index += 1

    for node in rewrite_nodes:
        for num_id in node.iter(num_id_tag):
            old_value = _integer_attribute(num_id, value_attr)
            if old_value in num_map:
                num_id.set(value_attr, str(num_map[old_value]))

    merger.target.parts[target_part] = _serialize_xml(target_root)
    merger.target_content_types.ensure(
        target_part,
        merger.source_content_types.get(source_part) or NUMBERING_CONTENT_TYPE,
    )
    merger.ensure_document_relationship(target_part, "/numbering")
    return target_root, len(copied_nums)


def _merge_note_part(
    merger: _OpcMerger,
    body_nodes: Sequence[etree._Element],
    word_ns: str,
    *,
    singular: str,
    plural: str,
    content_type: str,
) -> list[etree._Element]:
    reference_tag = _qn(word_ns, f"{singular}Reference")
    id_attr = _qn(word_ns, "id")
    reference_elements = [
        element
        for node in body_nodes
        for element in node.iter(reference_tag)
        if (_integer_attribute(element, id_attr) or 0) >= 0
    ]
    used_ids = {
        value
        for element in reference_elements
        for value in [_integer_attribute(element, id_attr)]
        if value is not None
    }
    if not used_ids:
        return []

    source_part = merger.source_related_part(f"/{plural}", f"word/{plural}.xml")
    target_part = merger.target_related_part(f"/{plural}", f"word/{plural}.xml")
    source_data = merger.source.parts.get(source_part)
    if source_data is None:
        # Same-reference-doc fallback: leave IDs alone if the template owns the
        # referenced note definitions.
        target_data = merger.target.parts.get(target_part)
        if target_data is None:
            raise TemplateFillError(f"Content uses {plural} but has no {plural}.xml")
        target_root = _parse_xml(target_data, target_part)
        target_ids = {
            _integer_attribute(element, id_attr)
            for element in target_root
            if _local_name(element.tag) == singular
        }
        missing = used_ids - {value for value in target_ids if value is not None}
        if missing:
            raise TemplateFillError(
                f"Content references missing {plural}: {sorted(missing)}"
            )
        return []

    source_root = _parse_xml(source_data, source_part)
    if target_part in merger.target.parts:
        target_root = _parse_xml(merger.target.parts[target_part], target_part)
    else:
        target_root = etree.Element(
            _qn(word_ns, plural), nsmap={"w": word_ns, "r": R_NS}
        )
        # Separator notes are package infrastructure rather than user content.
        for note in source_root:
            note_id = _integer_attribute(note, id_attr)
            if _local_name(note.tag) == singular and note_id is not None and note_id < 1:
                target_root.append(copy.deepcopy(note))

    source_notes = {
        _integer_attribute(element, id_attr): element
        for element in source_root
        if _local_name(element.tag) == singular
    }
    target_ids = {
        value
        for element in target_root
        if _local_name(element.tag) == singular
        for value in [_integer_attribute(element, id_attr)]
        if value is not None
    }
    next_id = max({value for value in target_ids if value > 0} or {0}) + 1
    id_map: dict[int, int] = {}
    copies: list[etree._Element] = []
    for old_id in sorted(used_ids):
        source_note = source_notes.get(old_id)
        if source_note is None:
            if old_id in target_ids:
                id_map[old_id] = old_id
                continue
            raise TemplateFillError(
                f"Content {plural}.xml has no {singular} id={old_id}"
            )
        while next_id in target_ids:
            next_id += 1
        new_id = next_id
        next_id += 1
        target_ids.add(new_id)
        copied = copy.deepcopy(source_note)
        copied.set(id_attr, str(new_id))
        target_root.append(copied)
        copies.append(copied)
        id_map[old_id] = new_id

    for reference in reference_elements:
        old_id = _integer_attribute(reference, id_attr)
        if old_id in id_map:
            reference.set(id_attr, str(id_map[old_id]))

    merger.merge_relationships_for_nodes(source_part, target_part, copies)
    merger.target.parts[target_part] = _serialize_xml(target_root)
    merger.target_content_types.ensure(
        target_part,
        merger.source_content_types.get(source_part) or content_type,
    )
    merger.ensure_document_relationship(target_part, f"/{plural}")
    return copies


def _merge_comments(
    merger: _OpcMerger,
    body_nodes: Sequence[etree._Element],
    word_ns: str,
) -> list[etree._Element]:
    marker_tags = {
        _qn(word_ns, "commentRangeStart"),
        _qn(word_ns, "commentRangeEnd"),
        _qn(word_ns, "commentReference"),
    }
    id_attr = _qn(word_ns, "id")
    markers = [
        element
        for node in body_nodes
        for element in node.iter()
        if element.tag in marker_tags
    ]
    used_ids = {
        value
        for marker in markers
        for value in [_integer_attribute(marker, id_attr)]
        if value is not None
    }
    if not used_ids:
        return []

    source_part = merger.source_related_part("/comments", "word/comments.xml")
    target_part = merger.target_related_part("/comments", "word/comments.xml")
    source_data = merger.source.parts.get(source_part)
    if source_data is None:
        raise TemplateFillError("Content uses comments but has no comments.xml")
    source_root = _parse_xml(source_data, source_part)
    if target_part in merger.target.parts:
        target_root = _parse_xml(merger.target.parts[target_part], target_part)
    else:
        target_root = etree.Element(
            _qn(word_ns, "comments"), nsmap={"w": word_ns, "r": R_NS}
        )

    source_comments = {
        _integer_attribute(element, id_attr): element
        for element in source_root
        if _local_name(element.tag) == "comment"
    }
    target_ids = {
        value
        for element in target_root
        if _local_name(element.tag) == "comment"
        for value in [_integer_attribute(element, id_attr)]
        if value is not None
    }
    next_id = max(target_ids or {-1}) + 1
    id_map: dict[int, int] = {}
    copies: list[etree._Element] = []
    for old_id in sorted(used_ids):
        source_comment = source_comments.get(old_id)
        if source_comment is None:
            raise TemplateFillError(f"comments.xml has no comment id={old_id}")
        while next_id in target_ids:
            next_id += 1
        copied = copy.deepcopy(source_comment)
        copied.set(id_attr, str(next_id))
        target_root.append(copied)
        copies.append(copied)
        id_map[old_id] = next_id
        target_ids.add(next_id)
        next_id += 1

    for marker in markers:
        old_id = _integer_attribute(marker, id_attr)
        if old_id in id_map:
            marker.set(id_attr, str(id_map[old_id]))

    merger.merge_relationships_for_nodes(source_part, target_part, copies)
    merger.target.parts[target_part] = _serialize_xml(target_root)
    merger.target_content_types.ensure(
        target_part,
        merger.source_content_types.get(source_part) or COMMENTS_CONTENT_TYPE,
    )
    merger.ensure_document_relationship(target_part, "/comments")
    return copies


def _renumber_drawing_properties(
    template_root: etree._Element,
    inserted_nodes: Sequence[etree._Element],
) -> None:
    for namespace, local_name in ((WP_NS, "docPr"), (PIC_NS, "cNvPr")):
        tag = _qn(namespace, local_name)
        existing_ids = {
            value
            for element in template_root.iter(tag)
            for value in [_integer_attribute(element, "id")]
            if value is not None
        }
        next_id = max(existing_ids or {0}) + 1
        for node in inserted_nodes:
            for element in node.iter(tag):
                while next_id in existing_ids:
                    next_id += 1
                element.set("id", str(next_id))
                existing_ids.add(next_id)
                next_id += 1


def _renumber_bookmarks(
    template_root: etree._Element,
    inserted_nodes: Sequence[etree._Element],
    word_ns: str,
) -> None:
    start_tag = _qn(word_ns, "bookmarkStart")
    end_tag = _qn(word_ns, "bookmarkEnd")
    hyperlink_tag = _qn(word_ns, "hyperlink")
    id_attr = _qn(word_ns, "id")
    name_attr = _qn(word_ns, "name")
    anchor_attr = _qn(word_ns, "anchor")
    existing_ids = {
        value
        for element in template_root.iter(start_tag)
        for value in [_integer_attribute(element, id_attr)]
        if value is not None
    }
    existing_names = {
        element.get(name_attr)
        for element in template_root.iter(start_tag)
        if element.get(name_attr)
    }
    next_id = max(existing_ids or {-1}) + 1
    id_map: dict[int, int] = {}
    name_map: dict[str, str] = {}

    for node in inserted_nodes:
        for start in node.iter(start_tag):
            old_id = _integer_attribute(start, id_attr)
            if old_id is not None:
                if old_id not in id_map:
                    while next_id in existing_ids:
                        next_id += 1
                    id_map[old_id] = next_id
                    existing_ids.add(next_id)
                    next_id += 1
                start.set(id_attr, str(id_map[old_id]))
            old_name = start.get(name_attr)
            if old_name and old_name in existing_names:
                candidate_base = f"content_{old_name.lstrip('_')}" or "content_bookmark"
                candidate = candidate_base
                suffix = 2
                while candidate in existing_names:
                    candidate = f"{candidate_base}_{suffix}"
                    suffix += 1
                name_map[old_name] = candidate
                start.set(name_attr, candidate)
                existing_names.add(candidate)
            elif old_name:
                existing_names.add(old_name)

    for node in inserted_nodes:
        for end in node.iter(end_tag):
            old_id = _integer_attribute(end, id_attr)
            if old_id in id_map:
                end.set(id_attr, str(id_map[old_id]))
        for hyperlink in node.iter(hyperlink_tag):
            old_anchor = hyperlink.get(anchor_attr)
            if old_anchor in name_map:
                hyperlink.set(anchor_attr, name_map[old_anchor])


def _anchor_requires_preservation(anchor: etree._Element, word_ns: str) -> bool:
    if anchor.find(f".//{_qn(word_ns, 'sectPr')}") is not None:
        return True
    if anchor.find(f".//{_qn(word_ns, 'pageBreakBefore')}") is not None:
        return True
    for line_break in anchor.iter(_qn(word_ns, "br")):
        if (line_break.get(_qn(word_ns, "type")) or "textWrapping") == "page":
            return True
    return False


def _clear_placeholder_text(anchor: etree._Element, word_ns: str) -> None:
    removable_tags = {
        _qn(word_ns, "t"),
        _qn(word_ns, "delText"),
        _qn(word_ns, "instrText"),
    }
    for element in list(anchor.iter()):
        if element.tag in removable_tags:
            parent = element.getparent()
            if parent is not None:
                parent.remove(element)


def _disable_content_control_placeholder(anchor: etree._Element, word_ns: str) -> None:
    current: etree._Element | None = anchor
    while current is not None:
        if current.tag == _qn(word_ns, "sdtContent"):
            sdt = current.getparent()
            if sdt is not None and sdt.tag == _qn(word_ns, "sdt"):
                sdt_pr = sdt.find(_qn(word_ns, "sdtPr"))
                if sdt_pr is not None:
                    for marker in list(sdt_pr.findall(_qn(word_ns, "showingPlcHdr"))):
                        sdt_pr.remove(marker)
            return
        current = current.getparent()


def _insert_at_anchor(
    anchor: etree._Element,
    nodes: Sequence[etree._Element],
    word_ns: str,
    anchor_mode: str,
) -> None:
    parent = anchor.getparent()
    if parent is None:
        raise TemplateFillError("Resolved body anchor has no parent element")
    mode = anchor_mode.strip().lower().replace("-", "_")
    if mode in {"insert_after", "after"}:
        insertion_index = parent.index(anchor) + 1
        remove_anchor = False
    elif mode in {"insert_before", "before", "keep"}:
        insertion_index = parent.index(anchor)
        remove_anchor = False
    elif mode in {"delete", "replace", "replace_anchor", ""}:
        if _anchor_requires_preservation(anchor, word_ns):
            # Section/page-break properties are part of the template layout.
            # Keep their paragraph, remove only placeholder text, and insert the
            # body after it.
            _clear_placeholder_text(anchor, word_ns)
            insertion_index = parent.index(anchor) + 1
            remove_anchor = False
        else:
            insertion_index = parent.index(anchor)
            remove_anchor = True
    else:
        raise TemplateFillError(
            f"Unsupported body_anchor_mode {anchor_mode!r}; expected delete, before, or after"
        )

    _disable_content_control_placeholder(anchor, word_ns)
    for node in nodes:
        parent.insert(insertion_index, node)
        insertion_index += 1
    if remove_anchor:
        parent.remove(anchor)


def _validate_input_path(path: Path, allowed_suffixes: set[str], label: str) -> Path:
    path = path.expanduser().resolve()
    if not path.is_file():
        raise TemplateFillError(f"{label} does not exist or is not a file: {str(path)!r}")
    if path.suffix.lower() not in allowed_suffixes:
        expected = ", ".join(sorted(allowed_suffixes))
        raise TemplateFillError(f"{label} must use one of: {expected}")
    return path


def _validate_generated_docx(output: Path) -> None:
    # python-docx gives a useful second validation layer on top of XML/ZIP
    # parsing and is already a project runtime dependency.
    try:
        from docx import Document

        Document(str(output))
    except ImportError:
        # lxml/zipfile validation above is sufficient for minimal installations.
        return
    except Exception as exc:  # pragma: no cover - implementation-specific detail
        raise TemplateFillError(f"Generated output is not a readable DOCX: {exc}") from exc


def fill_template(
    *,
    template: str | os.PathLike[str],
    content_docx: str | os.PathLike[str],
    output: str | os.PathLike[str],
    map_path: str | os.PathLike[str] | None = None,
) -> FillReport:
    """Fill *template* with the body elements from *content_docx*.

    The function is intentionally keyword-only so workflow integrations cannot
    accidentally swap the two DOCX inputs.
    """

    template_path = _validate_input_path(
        Path(template), {".docx", ".dotx"}, "Template"
    )
    content_path = _validate_input_path(
        Path(content_docx), {".docx"}, "Content DOCX"
    )
    output_path = Path(output).expanduser().resolve()
    if output_path.suffix.lower() != ".docx":
        raise TemplateFillError("Output path must end in .docx")
    resolved_map = (
        _validate_input_path(Path(map_path), {".json"}, "Template map")
        if map_path is not None
        else None
    )
    template_map = _load_template_map(resolved_map)

    template_package = _Package(template_path)
    content_package = _Package(content_path)
    template_document = _parse_xml(
        template_package.parts[DOCUMENT_PART], DOCUMENT_PART
    )
    content_document = _parse_xml(
        content_package.parts[DOCUMENT_PART], DOCUMENT_PART
    )
    template_word_ns = _namespace(template_document.tag)
    content_word_ns = _namespace(content_document.tag)
    if template_word_ns != content_word_ns:
        raise TemplateFillError(
            "Template and content use different WordprocessingML dialects; "
            "convert both to standard DOCX first"
        )
    word_ns = template_word_ns or W_NS
    template_body = _body_element(template_document, word_ns)
    content_body = _body_element(content_document, word_ns)

    anchor: etree._Element | None = None
    anchor_source: str | None = None
    map_error: TemplateFillError | None = None
    if template_map:
        try:
            anchor, anchor_source = _anchor_from_map(
                template_body, word_ns, template_map
            )
        except TemplateFillError as exc:
            map_error = exc
    if anchor is None:
        anchor, placeholder_source = _find_placeholder(template_body, word_ns)
        if anchor is not None:
            anchor_source = placeholder_source
    if anchor is None:
        if map_error is not None:
            raise TemplateFillError(
                f"Template-map anchor failed ({map_error}) and no common body placeholder was found"
            ) from map_error
        raise TemplateFillError(
            "No body anchor found. Provide --map with body_anchor_para_idx or add "
            "a common placeholder such as {{BODY}} / {{正文}} to the template."
        )

    original_anchor_text = _paragraph_text(anchor, word_ns)
    inserted_nodes = _content_body_nodes(content_body, word_ns)
    merger = _OpcMerger(template_package, content_package)

    # First import all package relationships directly used by body content
    # (images, hyperlinks, charts, embedded objects, altChunk parts, ...).
    merger.merge_relationships_for_nodes(
        DOCUMENT_PART, DOCUMENT_PART, inserted_nodes
    )

    # Notes/comments have their own relationship scopes and definitions.
    auxiliary_nodes: list[etree._Element] = []
    auxiliary_nodes.extend(
        _merge_note_part(
            merger,
            inserted_nodes,
            word_ns,
            singular="footnote",
            plural="footnotes",
            content_type=FOOTNOTES_CONTENT_TYPE,
        )
    )
    auxiliary_nodes.extend(
        _merge_note_part(
            merger,
            inserted_nodes,
            word_ns,
            singular="endnote",
            plural="endnotes",
            content_type=ENDNOTES_CONTENT_TYPE,
        )
    )
    auxiliary_nodes.extend(_merge_comments(merger, inserted_nodes, word_ns))

    all_inserted_content = [*inserted_nodes, *auxiliary_nodes]
    styles_root, copied_styles = _merge_styles(
        merger, all_inserted_content, word_ns
    )
    _, imported_numbering = _merge_numbering(
        merger, [*all_inserted_content, *copied_styles], word_ns
    )
    if styles_root is not None:
        styles_part = merger.target_related_part("/styles", "word/styles.xml")
        template_package.parts[styles_part] = _serialize_xml(styles_root)

    # Avoid duplicate drawing/bookmark IDs after two document trees are joined.
    _renumber_drawing_properties(template_document, inserted_nodes)
    _renumber_bookmarks(template_document, inserted_nodes, word_ns)

    anchor_mode = str(template_map.get("body_anchor_mode", "delete"))
    _insert_at_anchor(anchor, inserted_nodes, word_ns, anchor_mode)
    template_package.parts[DOCUMENT_PART] = _serialize_xml(template_document)

    merger.target_content_types.force_docx_main_document()
    merger.flush_relationships()
    merger.target_content_types.flush()
    template_package.write(output_path)
    _validate_generated_docx(output_path)

    return FillReport(
        template=str(template_path),
        content_docx=str(content_path),
        output=str(output_path),
        anchor_source=anchor_source or "unknown",
        anchor_text=original_anchor_text,
        inserted_elements=len(inserted_nodes),
        imported_relationships=merger.imported_relationships,
        imported_parts=merger.imported_parts,
        imported_styles=len(copied_styles),
        imported_numbering_instances=imported_numbering,
    )


def _build_argument_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description=(
            "Insert an already-rendered DOCX body into a DOCX/DOTX template "
            "without rebuilding the template package."
        )
    )
    parser.add_argument("--template", required=True, help="Input .docx or .dotx template")
    parser.add_argument(
        "--content-docx",
        required=True,
        help="Already-rendered .docx whose body will be inserted",
    )
    parser.add_argument("--output", required=True, help="Output .docx path")
    parser.add_argument(
        "--map",
        dest="map_path",
        help="Optional _template_map.json containing body_anchor_para_idx",
    )
    return parser


def main(argv: Sequence[str] | None = None) -> int:
    parser = _build_argument_parser()
    args = parser.parse_args(argv)
    try:
        report = fill_template(
            template=args.template,
            content_docx=args.content_docx,
            output=args.output,
            map_path=args.map_path,
        )
    except TemplateFillError as exc:
        print(f"docx-template-fill: error: {exc}", file=sys.stderr)
        return 2
    print(json.dumps(asdict(report), ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
